"""
Generadores de certificados PDF (reportlab) y Word (python-docx).
"""
from io import BytesIO
from datetime import datetime

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.colors import HexColor
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    import qrcode
    QR_DISPONIBLE = True
except ImportError:
    QR_DISPONIBLE = False


# -------------------------------------------------------------
# Colores corporativos
# -------------------------------------------------------------
AZUL = HexColor("#1e3a8a")
ROJO = HexColor("#b91c1c")
DORADO = HexColor("#ca8a04")
GRIS = HexColor("#4b5563")
FONDO = HexColor("#fefce8")


# -------------------------------------------------------------
# PDF con diseño diploma
# -------------------------------------------------------------
def generar_certificado_pdf(certificado, alumno, firma1="Profesor", firma2="Dirección Académica"):
    """Genera un PDF tipo diploma y devuelve los bytes."""
    buffer = BytesIO()
    page = landscape(A4)  # 29.7 x 21 cm
    W, H = page
    p = canvas.Canvas(buffer, pagesize=page)

    # -------- Fondo --------
    p.setFillColor(FONDO)
    p.rect(0, 0, W, H, fill=1, stroke=0)

    # -------- Bordes decorativos --------
    p.setStrokeColor(AZUL)
    p.setLineWidth(8)
    p.rect(1*cm, 1*cm, W-2*cm, H-2*cm, fill=0, stroke=1)

    p.setStrokeColor(DORADO)
    p.setLineWidth(3)
    p.rect(1.4*cm, 1.4*cm, W-2.8*cm, H-2.8*cm, fill=0, stroke=1)

    p.setStrokeColor(AZUL)
    p.setLineWidth(1)
    p.rect(1.7*cm, 1.7*cm, W-3.4*cm, H-3.4*cm, fill=0, stroke=1)

    # -------- Esquinas decorativas --------
    p.setFillColor(DORADO)
    for x, y in [(1.7*cm, 1.7*cm), (W-1.7*cm, 1.7*cm), (1.7*cm, H-1.7*cm), (W-1.7*cm, H-1.7*cm)]:
        p.circle(x, y, 0.25*cm, fill=1, stroke=0)

    # -------- Encabezado --------
    p.setFillColor(AZUL)
    p.setFont("Helvetica-Bold", 14)
    p.drawCentredString(W/2, H - 2.8*cm, "RHETORICA+")

    p.setFillColor(GRIS)
    p.setFont("Helvetica-Oblique", 10)
    p.drawCentredString(W/2, H - 3.4*cm, "Plataforma de aprendizaje retórico basada en IA")

    # Línea decorativa bajo el encabezado
    p.setStrokeColor(DORADO)
    p.setLineWidth(2)
    p.line(W/2 - 4*cm, H - 3.8*cm, W/2 + 4*cm, H - 3.8*cm)

    # -------- Título --------
    p.setFillColor(AZUL)
    p.setFont("Helvetica-Bold", 44)
    p.drawCentredString(W/2, H - 5.5*cm, "CERTIFICADO")

    p.setFillColor(DORADO)
    p.setFont("Helvetica-Bold", 14)
    p.drawCentredString(W/2, H - 6.4*cm, "DE FINALIZACIÓN")

    # -------- Cuerpo --------
    p.setFillColor(GRIS)
    p.setFont("Helvetica", 12)
    p.drawCentredString(W/2, H - 7.6*cm, "Se otorga el presente certificado a:")

    # Nombre del alumno
    p.setFillColor(AZUL)
    p.setFont("Helvetica-BoldOblique", 32)
    nombre_completo = alumno.get_full_name() or alumno.username
    p.drawCentredString(W/2, H - 9.2*cm, nombre_completo)

    # Línea bajo el nombre
    p.setStrokeColor(GRIS)
    p.setLineWidth(0.5)
    ancho_nombre = p.stringWidth(nombre_completo, "Helvetica-BoldOblique", 32)
    p.line(W/2 - ancho_nombre/2 - 1*cm, H - 9.6*cm, W/2 + ancho_nombre/2 + 1*cm, H - 9.6*cm)

    # Texto
    p.setFillColor(GRIS)
    p.setFont("Helvetica", 12)
    p.drawCentredString(W/2, H - 10.6*cm, "por haber completado satisfactoriamente la certificación:")

    # Nombre de certificación
    p.setFillColor(ROJO)
    p.setFont("Helvetica-Bold", 20)
    p.drawCentredString(W/2, H - 11.8*cm, certificado.certificacion.nombre)

    # Nivel
    p.setFillColor(GRIS)
    p.setFont("Helvetica-Oblique", 11)
    p.drawCentredString(W/2, H - 12.6*cm, f"Nivel: {certificado.certificacion.get_nivel_display() if hasattr(certificado.certificacion, 'get_nivel_display') else certificado.certificacion.nivel.capitalize()}")

    # -------- Firmas --------
    y_firma = H - 15.5*cm
    ancho_firma = 6*cm

    # Firma izquierda
    x1 = W/2 - ancho_firma - 1.5*cm
    p.setStrokeColor(GRIS)
    p.setLineWidth(0.8)
    p.line(x1, y_firma, x1 + ancho_firma, y_firma)
    p.setFillColor(GRIS)
    p.setFont("Helvetica-Bold", 10)
    p.drawCentredString(x1 + ancho_firma/2, y_firma - 0.5*cm, firma1)
    p.setFont("Helvetica", 8)
    p.setFillColor(HexColor("#9ca3af"))
    p.drawCentredString(x1 + ancho_firma/2, y_firma - 0.95*cm, "Firma autorizada")

    # Firma derecha
    x2 = W/2 + 1.5*cm
    p.setStrokeColor(GRIS)
    p.line(x2, y_firma, x2 + ancho_firma, y_firma)
    p.setFillColor(GRIS)
    p.setFont("Helvetica-Bold", 10)
    p.drawCentredString(x2 + ancho_firma/2, y_firma - 0.5*cm, firma2)
    p.setFont("Helvetica", 8)
    p.setFillColor(HexColor("#9ca3af"))
    p.drawCentredString(x2 + ancho_firma/2, y_firma - 0.95*cm, "Firma autorizada")

    # -------- Sello central --------
    p.setStrokeColor(DORADO)
    p.setLineWidth(1.5)
    p.circle(W/2, y_firma + 0.4*cm, 0.9*cm, fill=0, stroke=1)
    p.setFillColor(DORADO)
    p.setFont("Helvetica-Bold", 7)
    p.drawCentredString(W/2, y_firma + 0.5*cm, "SELLO")
    p.drawCentredString(W/2, y_firma + 0.15*cm, "OFICIAL")

    # -------- Fecha y código --------
    y_meta = 3.2*cm
    p.setFillColor(GRIS)
    p.setFont("Helvetica", 9)
    fecha_str = certificado.fecha_obtenida.strftime("%d de %B de %Y")
    meses_es = {
        "January": "enero", "February": "febrero", "March": "marzo",
        "April": "abril", "May": "mayo", "June": "junio",
        "July": "julio", "August": "agosto", "September": "septiembre",
        "October": "octubre", "November": "noviembre", "December": "diciembre",
    }
    for en, es in meses_es.items():
        fecha_str = fecha_str.replace(en, es)

    p.drawString(2.5*cm, y_meta, f"Fecha de emisión: {fecha_str}")
    p.drawString(2.5*cm, y_meta - 0.5*cm, f"Código de verificación: {certificado.codigo_verificacion}")

    url_verif = certificado.url_verificacion or f"https://rhetorica.com/verificar/{certificado.codigo_verificacion}"
    p.drawString(2.5*cm, y_meta - 1*cm, f"Verificar en: {url_verif}")

    # -------- QR --------
    if QR_DISPONIBLE:
        try:
            qr = qrcode.QRCode(box_size=3, border=1)
            qr.add_data(url_verif)
            qr.make(fit=True)
            img = qr.make_image(fill_color="#1e3a8a", back_color="#fefce8")
            qr_buffer = BytesIO()
            img.save(qr_buffer, format="PNG")
            qr_buffer.seek(0)
            from reportlab.lib.utils import ImageReader
            p.drawImage(ImageReader(qr_buffer), W - 4*cm, y_meta - 0.5*cm, 2.5*cm, 2.5*cm, mask='auto')
        except Exception:
            pass

    # -------- Footer --------
    p.setFillColor(HexColor("#9ca3af"))
    p.setFont("Helvetica-Oblique", 7)
    p.drawCentredString(W/2, 0.8*cm, "Este certificado puede verificarse en línea usando el código de verificación.")

    p.showPage()
    p.save()
    return buffer.getvalue()


# -------------------------------------------------------------
# Word editable
# -------------------------------------------------------------
def generar_certificado_docx(certificado, alumno, firma1="Profesor", firma2="Dirección Académica"):
    """Genera un .docx editable y devuelve los bytes."""
    doc = Document()

    # Configurar márgenes (A4 landscape)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.page_width = Cm(29.7)
        section.page_height = Cm(21)

    # Encabezado
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("RHETORICA+")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Plataforma de aprendizaje retórico basada en IA")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    # Título
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(20)
    r = t.add_run("CERTIFICADO")
    r.bold = True
    r.font.size = Pt(44)
    r.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run("DE FINALIZACIÓN")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0xca, 0x8a, 0x04)

    # Cuerpo
    doc.add_paragraph()
    c1 = doc.add_paragraph()
    c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c1.add_run("Se otorga el presente certificado a:")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    # Nombre
    n = doc.add_paragraph()
    n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = n.add_run(alumno.get_full_name() or alumno.username)
    r.bold = True
    r.italic = True
    r.font.size = Pt(32)
    r.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)

    # Texto
    c2 = doc.add_paragraph()
    c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c2.add_run("por haber completado satisfactoriamente la certificación:")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    # Nombre de certificación
    cn = doc.add_paragraph()
    cn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cn.add_run(certificado.certificacion.nombre)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)

    # Nivel
    nv = doc.add_paragraph()
    nv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = nv.add_run(f"Nivel: {certificado.certificacion.nivel.capitalize()}")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    # Firmas (tabla 2 columnas)
    doc.add_paragraph()
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(8)
    table.columns[1].width = Cm(8)

    # Firma 1
    cell1 = table.cell(0, 0)
    p1 = cell1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1.add_run("_________________________")
    p1_2 = cell1.add_paragraph()
    p1_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1_2.add_run(firma1)
    r.bold = True
    r.font.size = Pt(10)

    # Firma 2
    cell2 = table.cell(0, 1)
    p2 = cell2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("_________________________")
    p2_2 = cell2.add_paragraph()
    p2_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2_2.add_run(firma2)
    r.bold = True
    r.font.size = Pt(10)

    # Fecha y código
    doc.add_paragraph()
    doc.add_paragraph()
    fecha = certificado.fecha_obtenida.strftime("%d/%m/%Y")
    f = doc.add_paragraph()
    r = f.add_run(f"Fecha de emisión: {fecha}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    cod = doc.add_paragraph()
    r = cod.add_run(f"Código de verificación: {certificado.codigo_verificacion}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    fol = doc.add_paragraph()
    r = fol.add_run(f"Folio: {folio_unico(certificado)}   |   Huella: {huella_corta(certificado)}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    url = certificado.url_verificacion or f"https://rhetorica.com/verificar/{certificado.codigo_verificacion}"
    u = doc.add_paragraph()
    r = u.add_run(f"Verificar en: {url}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x4b, 0x55, 0x63)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------
# Huella / hash de integridad del certificado
# -------------------------------------------------------------
import hashlib


def huella_certificado(certificado):
    """Genera un hash SHA-256 determinista a partir de los datos del certificado."""
    partes = [
        str(certificado.id),
        certificado.codigo_verificacion or "",
        certificado.user.username or "",
        certificado.certificacion.nombre or "",
        certificado.fecha_obtenida.strftime("%Y-%m-%d") if certificado.fecha_obtenida else "",
    ]
    base = "|".join(partes).encode("utf-8")
    return hashlib.sha256(base).hexdigest().upper()


def huella_corta(certificado):
    """Versión corta y legible: XXXX-XXXX-XXXX-XXXX"""
    h = huella_certificado(certificado)
    return "-".join([h[i:i+4] for i in range(0, 16, 4)])


def folio_unico(certificado):
    """Folio tipo RH-2026-XXXXXX"""
    anio = certificado.fecha_obtenida.year if certificado.fecha_obtenida else "0000"
    return f"RH-{anio}-{certificado.id:06d}"


def url_publica(certificado):
    """URL pública para verificar (funciona sin login)."""
    return f"/verificar/{certificado.codigo_verificacion}/"
